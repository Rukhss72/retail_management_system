from django.core.management.base import BaseCommand
from django.db import transaction
from retail.models import FestiveCalendar

from openpyxl import load_workbook
from datetime import datetime, date


def parse_date(value):
    """
    Accepts:
    - Excel date (datetime/date)
    - string: dd/mm/yyyy OR yyyy-mm-dd
    Returns: date or None
    """
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass

    raise ValueError(f"Unrecognized date value: {value!r}")


class Command(BaseCommand):
    help = "Import FestiveCalendar records from an XLSX file (dd/mm/yyyy supported)."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str, help="C:/OneDrive to disk/Ruk IOT/CAPSTONE Project/Retail Management System/festive_calendar_2023_2027.xlsx")
        parser.add_argument("--sheet", type=str, default=None, help="Sheet name (default: active sheet)")
        parser.add_argument("--clear", action="store_true", help="Delete existing festive_calendar before importing")

    @transaction.atomic
    def handle(self, *args, **options):
        xlsx_path = options["xlsx_path"]
        sheet_name = options["sheet"]
        clear = options["clear"]

        wb = load_workbook(filename=xlsx_path, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active

        # Expect headers in row 1
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        header_map = {h: idx for idx, h in enumerate(headers)}

        required = ["Festival_name", "Start_date", "End_date", "Festival_date"]
        for r in required:
            if r not in header_map:
                raise ValueError(f"Missing required column: {r}. Found: {headers}")

        if clear:
            FestiveCalendar.objects.all().delete()
            self.stdout.write(self.style.WARNING("Cleared existing FestiveCalendar rows."))

        created, updated = 0, 0

        # rows from row 2 onward
        for row in ws.iter_rows(min_row=2, values_only=True):
            festival_name = row[header_map["Festival_name"]]
            if not festival_name:
                continue

            festival_name = str(festival_name).strip()
            start_date = parse_date(row[header_map["Start_date"]])
            end_date = parse_date(row[header_map["End_date"]])
            festival_date = parse_date(row[header_map["Festival_date"]])

            obj, was_created = FestiveCalendar.objects.update_or_create(
                festival_name=festival_name,
                start_date=start_date,
                end_date=end_date,
                festival_date=festival_date,
                defaults={},  # nothing extra for now
            )
            created += 1 if was_created else 0
            updated += 0 if was_created else 1

        self.stdout.write(self.style.SUCCESS(
            f"Import done. Created: {created}, Updated: {updated}"
        ))
